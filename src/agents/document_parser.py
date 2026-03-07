"""
文档解析Agent
负责解析招标文件（docx/pdf），提取结构化内容
"""

import re
from pathlib import Path
from typing import Dict, List, Any, Optional
from dataclasses import dataclass, field
import json

try:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    Document = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


@dataclass
class Section:
    """章节信息"""
    level: int
    title: str
    content: str
    page_start: int
    page_end: int
    subsections: List['Section'] = field(default_factory=list)


@dataclass
class TableData:
    """表格数据"""
    rows: List[List[str]]
    page_num: int
    headers: List[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    """解析后的文档结构"""
    filename: str
    file_type: str
    total_pages: int
    sections: List[Section]
    paragraphs: List[Dict[str, Any]]  # {text, page_num, section}
    tables: List[TableData]
    attachments: List[Dict[str, Any]] = field(default_factory=list)


class DocumentParserAgent:
    """文档解析Agent"""

    def __init__(self):
        self.supported_types = ['.docx', '.doc', '.pdf']

    def parse(self, file_path: str) -> ParsedDocument:
        """
        解析文档

        Args:
            file_path: 文件路径

        Returns:
            ParsedDocument: 解析后的文档结构
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        file_type = file_path.suffix.lower()
        if file_type not in self.supported_types:
            raise ValueError(f"不支持的文件类型: {file_type}")

        if file_type in ['.docx', '.doc']:
            return self._parse_docx(file_path)
        elif file_type == '.pdf':
            return self._parse_pdf(file_path)

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """解析Word文档"""
        if Document is None:
            raise ImportError("需要安装 python-docx: pip install python-docx")

        doc = Document(str(file_path))

        # 提取段落
        paragraphs = []
        current_page = 1

        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                paragraphs.append({
                    'text': para.text.strip(),
                    'page_num': current_page,
                    'section': self._infer_section(para),
                    'style': para.style.name if para.style else 'Normal'
                })

                # 简单估算页码（实际需要更复杂的计算）
                if i % 40 == 0 and i > 0:
                    current_page += 1

        # 提取表格
        tables = []
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)

            if table_data:
                tables.append(TableData(
                    rows=table_data,
                    page_num=current_page,
                    headers=table_data[0] if table_data else []
                ))

        # 构建章节结构
        sections = self._build_sections(paragraphs)

        return ParsedDocument(
            filename=file_path.name,
            file_type='docx',
            total_pages=current_page,
            sections=sections,
            paragraphs=paragraphs,
            tables=tables
        )

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """解析PDF文档"""
        if pdfplumber is None:
            raise ImportError("需要安装 pdfplumber: pip install pdfplumber")

        paragraphs = []
        tables = []
        current_page = 0

        with pdfplumber.open(str(file_path)) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                current_page = page_num
                text = page.extract_text() or ""

                # 按段落分割
                for line in text.split('\n'):
                    line = line.strip()
                    if line:
                        paragraphs.append({
                            'text': line,
                            'page_num': page_num,
                            'section': self._infer_section_from_text(line),
                            'style': 'Normal'
                        })

                # 提取表格
                page_tables = page.extract_tables()
                if page_tables:
                    for table_data in page_tables:
                        if table_data:
                            cleaned_table = [
                                [str(cell).strip() if cell else "" for cell in row]
                                for row in table_data
                            ]
                            tables.append(TableData(
                                rows=cleaned_table,
                                page_num=page_num,
                                headers=cleaned_table[0] if cleaned_table else []
                            ))

        sections = self._build_sections(paragraphs)

        return ParsedDocument(
            filename=file_path.name,
            file_type='pdf',
            total_pages=current_page,
            sections=sections,
            paragraphs=paragraphs,
            tables=tables
        )

    def _infer_section(self, paragraph: Paragraph) -> str:
        """从段落推断所属章节"""
        text = paragraph.text.strip()
        # 简单实现：检查是否为标题样式
        if paragraph.style and 'Heading' in paragraph.style.name:
            return text
        return ""

    def _infer_section_from_text(self, text: str) -> str:
        """从文本推断是否为章节标题"""
        # 检查常见章节标题模式
        patterns = [
            r'^第[一二三四五六七八九十]+[章节编]',
            r'^\d+\.\s',
            r'^[一二三四五六七八九十]+[、．]',
            r'^[一二三四五六七八九十]+、\w+',
        ]
        for pattern in patterns:
            if re.match(pattern, text):
                return text
        return ""

    def _build_sections(self, paragraphs: List[Dict[str, Any]]) -> List[Section]:
        """构建章节树结构"""
        sections = []
        section_stack = []

        for para in paragraphs:
            section_title = para.get('section', '')
            if not section_title:
                continue

            # 判断标题级别
            level = self._get_heading_level(section_title)

            # 创建新章节
            new_section = Section(
                level=level,
                title=section_title,
                content="",
                page_start=para['page_num'],
                page_end=para['page_num']
            )

            # 调整栈结构
            while section_stack and section_stack[-1].level >= level:
                section_stack.pop()

            if section_stack:
                section_stack[-1].subsections.append(new_section)
            else:
                sections.append(new_section)

            section_stack.append(new_section)

        return sections

    def _get_heading_level(self, title: str) -> int:
        """判断标题级别"""
        # 简单实现
        if re.match(r'^第.+[篇章编]', title):
            return 1
        elif re.match(r'^[一二三四五六七八九十]+[、．]', title):
            return 2
        elif re.match(r'^\d+\.\s', title):
            return 3
        return 4

    def extract_text_by_page(self, parsed_doc: ParsedDocument, page_num: int) -> str:
        """提取指定页码的文本"""
        texts = []
        for para in parsed_doc.paragraphs:
            if para['page_num'] == page_num:
                texts.append(para['text'])
        return '\n'.join(texts)

    def extract_text_by_section(self, parsed_doc: ParsedDocument, section_title: str) -> str:
        """提取指定章节的文本"""
        texts = []
        in_section = False

        for para in parsed_doc.paragraphs:
            if para.get('section') == section_title:
                in_section = True
            elif in_section and para.get('section'):
                # 遇到新章节，停止
                break
            elif in_section:
                texts.append(para['text'])

        return '\n'.join(texts)

    def search_by_keyword(self, parsed_doc: ParsedDocument, keyword: str) -> List[Dict[str, Any]]:
        """按关键词搜索"""
        results = []
        for para in parsed_doc.paragraphs:
            if keyword in para['text']:
                results.append({
                    'text': para['text'],
                    'page_num': para['page_num'],
                    'section': para.get('section', '')
                })
        return results

    def to_dict(self, parsed_doc: ParsedDocument) -> Dict[str, Any]:
        """转换为字典格式（用于存储）"""
        return {
            'filename': parsed_doc.filename,
            'file_type': parsed_doc.file_type,
            'total_pages': parsed_doc.total_pages,
            'sections': [
                {
                    'level': s.level,
                    'title': s.title,
                    'content': s.content,
                    'page_start': s.page_start,
                    'page_end': s.page_end
                }
                for s in parsed_doc.sections
            ],
            'paragraphs': parsed_doc.paragraphs,
            'tables': [
                {
                    'rows': t.rows,
                    'page_num': t.page_num,
                    'headers': t.headers
                }
                for t in parsed_doc.tables
            ]
        }

    def save_to_json(self, parsed_doc: ParsedDocument, output_path: str):
        """保存为JSON文件"""
        data = self.to_dict(parsed_doc)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)


# 使用示例
if __name__ == "__main__":
    parser = DocumentParserAgent()

    # 解析Word文档
    # doc = parser.parse("招标文件.docx")
    # parser.save_to_json(doc, "output.json")
