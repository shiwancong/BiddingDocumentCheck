#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
招标文件检查清单 Excel 生成器 v4.1
将 Markdown 格式的检查清单转换为 Excel 文件

结构说明（根据PDF官方要求调整）：
- 废标风险检查表：单个Sheet，按类型分区分组显示
- 建议项清单：单个Sheet，包含所有建议性要求
- 统计摘要：详细的总数和分类统计
- S级验证报告：质量验证信息

移除了"必须立即整改项"功能，保持与官方要求一致。
"""

import re
import json
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("警告: openpyxl 未安装，请运行: pip install openpyxl")


class ExcelGenerator:
    """Excel 文件生成器 v4.1"""

    # 风险等级权重，用于排序
    RISK_LEVEL_ORDER = {'高': 0, '中': 1, '低': 2}

    # 废标条款类型顺序
    CLAUSE_TYPES = ['资格', '文件', '签字盖章', '时间方式', '报价算术', '响应性', '格式', '模板', '其他']

    def __init__(self):
        self.wb = None
        self.data = {
            'metadata': {},
            'disqualification': [],      # 废标风险检查表数据
            'suggestions': [],           # 建议项清单数据
            'suggestion_legacy': [],     # 传统格式建议项（兼容）
            'statistics': {},            # 统计摘要
            'validation': {}             # 验证报告
        }

    def parse_markdown_file(self, md_file: str) -> bool:
        """解析 Markdown 格式的检查清单文件"""
        try:
            with open(md_file, 'r', encoding='utf-8') as f:
                content = f.read()
            return self._parse_content(content)
        except Exception as e:
            print(f"解析文件失败: {e}")
            return False

    def parse_json_data(self, json_file: str) -> bool:
        """解析 JSON 格式的数据文件"""
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                self.data = json.load(f)
            return True
        except Exception as e:
            print(f"解析 JSON 文件失败: {e}")
            return False

    def _parse_content(self, content: str) -> bool:
        """解析 Markdown 内容"""
        lines = content.split('\n')
        current_section = None
        current_category = None
        headers = []

        for line in lines:
            # 解析元数据
            if line.startswith('**生成时间**'):
                self.data['metadata']['generate_time'] = line.split(':', 1)[1].strip()
            elif line.startswith('**文件名称**'):
                self.data['metadata']['filename'] = line.split(':', 1)[1].strip()
            elif line.startswith('**总页数**'):
                self.data['metadata']['total_pages'] = line.split(':')[1].strip()
            elif line.startswith('**质量标准**'):
                self.data['metadata']['quality_level'] = line.split(':')[1].strip()

            # 检测当前所在表格
            if '## 废标风险检查表' in line:
                current_section = 'disqualification'
                current_category = None
            elif '## 建议项清单表' in line or '## 建议项清单' in line:
                current_section = 'suggestions'
                current_category = None
            elif '#### 表1：加分项建议' in line or '### 加分项建议' in line:
                current_category = '加分项'
            elif '#### 表2：优化项建议' in line or '### 优化项建议' in line:
                current_category = '优化项'
            elif '#### 表3：避坑建议' in line or '### 避坑建议' in line:
                current_category = '避坑项'

            # 解析表头
            if line.startswith('|') and ('检查项' in line or '建议项' in line):
                headers = [h.strip() for h in line.split('|')[1:-1]]

            # 解析表格行
            if line.startswith('|') and current_section in ['disqualification', 'suggestions']:
                # 跳过表头和分隔行
                if any(x in line for x in ['检查项', '建议项', '---', '风险等级', '加分潜力', '优先级', '参考条款']):
                    continue

                cells = [c.strip() for c in line.split('|')[1:-1]]
                if len(cells) < 2 or not cells[0]:
                    continue

                # 根据不同表格解析数据
                if current_section == 'disqualification':
                    item = self._parse_disqualification_row(cells, headers)
                    if item:
                        self.data['disqualification'].append(item)
                elif current_section == 'suggestions':
                    item = self._parse_suggestion_row(cells, headers, current_category)
                    if item:
                        self.data['suggestions'].append(item)

        # 按类型和风险等级排序废标条款
        self._sort_disqualification()

        return True

    def _parse_disqualification_row(self, cells: List[str], headers: List[str]) -> Dict:
        """解析废标风险检查表行数据"""
        # 检测是否有风险等级列
        has_risk_level = '风险等级' in headers

        if has_risk_level and len(cells) >= 9:
            # 新格式：风险等级 | 检查项 | 类型 | 原文页码 | 原文关键信息 | 逻辑解释 | 后果 | 参考页码 | 额外说明
            return {
                '风险等级': cells[0],
                '检查项': cells[1],
                '类型': cells[2],
                '原文页码': cells[3],
                '原文关键信息': cells[4],
                '逻辑解释': cells[5],
                '后果': cells[6],
                '参考页码': cells[7],
                '额外说明': cells[8] if len(cells) > 8 else ''
            }
        elif len(cells) >= 7:
            # 旧格式（兼容）：检查项 | 类型 | 原文页码 | 原文关键信息 | 逻辑解释 | 参考页码 | 额外说明
            risk_level = self._assess_risk_level(cells[4])
            return {
                '风险等级': risk_level,
                '检查项': cells[0],
                '类型': cells[1],
                '原文页码': cells[2],
                '原文关键信息': cells[3],
                '逻辑解释': cells[4],
                '后果': '',
                '参考页码': cells[5],
                '额外说明': cells[6] if len(cells) > 6 else ''
            }
        return {}

    def _parse_suggestion_row(self, cells: List[str], headers: List[str], category: str) -> Dict:
        """解析建议项清单表行数据"""
        # 检测表头格式
        has_new_format = any(x in headers for x in ['加分潜力', '优先级', '风险等级', '参考条款', '对应条款', '具体做法', '预期效果'])

        # 三类建议项新格式
        if has_new_format and len(cells) >= 7:
            item = {
                '类别': category or '其他',
                '序号': cells[0] if len(cells) > 0 else '',
                '建议项': cells[1] if len(cells) > 1 else '',
                '分类': cells[2] if len(cells) > 2 else '',
                '对应条款': cells[3] if len(cells) > 3 else '',
                '具体做法': cells[4] if len(cells) > 4 else '',
                '预期效果': cells[5] if len(cells) > 5 else '',
            }

            # 添加特定字段
            if len(cells) > 6:
                if category == '加分项':
                    item['加分潜力'] = cells[6]
                elif category == '优化项':
                    item['优先级'] = cells[6]
                elif category == '避坑项':
                    item['风险等级'] = cells[6]

            # 添加参考条款
            if len(cells) > 7:
                item['参考条款'] = cells[7]

            return item

        # 传统格式（兼容）
        elif len(cells) >= 7:
            return {
                '类别': '建议',
                '检查项': cells[0] if len(cells) > 0 else '',
                '类型': cells[1] if len(cells) > 1 else '',
                '原文页码': cells[2] if len(cells) > 2 else '',
                '原文关键信息': cells[3] if len(cells) > 3 else '',
                '逻辑解释': cells[4] if len(cells) > 4 else '',
                '参考页码': cells[5] if len(cells) > 5 else '',
                '额外说明': cells[6] if len(cells) > 6 else '',
                'legacy': True
            }

        return {}

    def _assess_risk_level(self, text: str) -> str:
        """根据文本内容自动评估风险等级"""
        text_lower = text.lower()

        # 高风险关键词
        high_keywords = ['否决', '无效', '不接受', '拒绝', '废标', '退回', '不得', '严禁']
        # 中风险关键词
        mid_keywords = ['视为', '按...处理', '可能', '应当', '必须']

        for keyword in high_keywords:
            if keyword in text_lower:
                return '高'

        for keyword in mid_keywords:
            if keyword in text_lower:
                return '中'

        return '低'

    def _sort_disqualification(self):
        """按类型和风险等级排序废标条款"""
        self.data['disqualification'].sort(
            key=lambda x: (
                self.CLAUSE_TYPES.index(x.get('类型', '其他')) if x.get('类型') in self.CLAUSE_TYPES else 999,
                self.RISK_LEVEL_ORDER.get(x.get('风险等级', '中'), 1),
                x.get('原文页码', '')
            )
        )

    def set_data(self, data: Dict[str, Any]):
        """直接设置数据"""
        self.data = data
        self._sort_disqualification()

    def generate_excel(self, output_file: str) -> bool:
        """生成 Excel 文件"""
        if not EXCEL_AVAILABLE:
            print("错误: openpyxl 未安装")
            return False

        try:
            self.wb = Workbook()

            # 删除默认 Sheet
            if 'Sheet' in self.wb.sheetnames:
                self.wb.remove(self.wb['Sheet'])

            # 创建各个 Sheet（按顺序）
            self._create_summary_sheet()
            self._create_disqualification_sheet()
            self._create_suggestions_sheet()
            self._create_validation_sheet()

            # 保存文件
            self.wb.save(output_file)
            print(f"[OK] Excel 文件已生成: {output_file}")

            # 输出统计信息
            self._print_statistics()

            return True

        except Exception as e:
            print(f"生成 Excel 失败: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _print_statistics(self):
        """打印统计信息"""
        disq = self.data.get('disqualification', [])
        sugg = self.data.get('suggestions', [])

        # 统计废标风险
        high_count = sum(1 for item in disq if item.get('风险等级') == '高')
        mid_count = sum(1 for item in disq if item.get('风险等级') == '中')
        low_count = sum(1 for item in disq if item.get('风险等级') == '低')

        # 统计建议项
        benefit_count = sum(1 for item in sugg if item.get('类别') == '加分项')
        optimize_count = sum(1 for item in sugg if item.get('类别') == '优化项')
        avoid_count = sum(1 for item in sugg if item.get('类别') == '避坑项')

        print(f"\n{'='*50}")
        print(f"  招投标文件智能检查清单 - 统计摘要")
        print(f"{'='*50}")
        print(f"\n[废标风险检查]")
        print(f"  总计: {len(disq)} 项")
        print(f"    高风险: {high_count} 项")
        print(f"    中风险: {mid_count} 项")
        print(f"    低风险: {low_count} 项")

        # 按类型统计
        type_stats = {}
        for item in disq:
            t = item.get('类型', '其他')
            type_stats[t] = type_stats.get(t, 0) + 1
        if type_stats:
            print(f"\n  按类型分布:")
            for t in self.CLAUSE_TYPES:
                if t in type_stats:
                    print(f"    {t}: {type_stats[t]} 项")

        print(f"\n[建议项清单]")
        print(f"  总计: {len(sugg)} 项")
        print(f"    加分项: {benefit_count} 项")
        print(f"    优化项: {optimize_count} 项")
        print(f"    避坑项: {avoid_count} 项")
        print(f"{'='*50}\n")

    def _create_summary_sheet(self):
        """创建统计摘要 Sheet"""
        ws = self.wb.create_sheet('统计摘要', 0)

        # 标题样式
        title_font = Font(size=14, bold=True, color='FFFFFF')
        title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        title_alignment = Alignment(horizontal='center', vertical='center')

        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 20

        # 写入主标题
        ws['A1'] = '招投标文件智能检查清单'
        ws['A1'].font = title_font
        ws['A1'].fill = title_fill
        ws['A1'].alignment = title_alignment
        ws.merge_cells('A1:C1')

        row = 3

        # 写入元数据
        if self.data.get('metadata'):
            ws.cell(row, 1, '基本信息').font = Font(bold=True, size=12)
            row += 1
            for key, value in self.data['metadata'].items():
                ws.cell(row, 1, key)
                ws.cell(row, 2, value)
                ws.cell(row, 1).font = Font(bold=True)
                row += 1
            row += 1

        # 写入废标风险统计
        disq = self.data.get('disqualification', [])
        if disq:
            ws.cell(row, 1, '废标风险统计').font = Font(bold=True, size=12, color='C00000')
            row += 1

            # 总计
            ws.cell(row, 1, '总计')
            ws.cell(row, 2, f'{len(disq)} 项')
            ws.cell(row, 2).font = Font(bold=True, size=14, color='C00000')
            row += 1

            # 按风险等级统计
            high_count = sum(1 for item in disq if item.get('风险等级') == '高')
            mid_count = sum(1 for item in disq if item.get('风险等级') == '中')
            low_count = sum(1 for item in disq if item.get('风险等级') == '低')

            ws.cell(row, 1, '  高风险')
            ws.cell(row, 2, f'{high_count} 项')
            ws.cell(row, 2).font = Font(color='C00000')
            row += 1

            ws.cell(row, 1, '  中风险')
            ws.cell(row, 2, f'{mid_count} 项')
            ws.cell(row, 2).font = Font(color='FF9900')
            row += 1

            ws.cell(row, 1, '  低风险')
            ws.cell(row, 2, f'{low_count} 项')
            ws.cell(row, 2).font = Font(color='008000')
            row += 2

            # 按类型统计
            type_stats = {}
            for item in disq:
                t = item.get('类型', '其他')
                type_stats[t] = type_stats.get(t, 0) + 1

            if type_stats:
                ws.cell(row, 1, '按类型分布').font = Font(bold=True)
                row += 1
                for t in self.CLAUSE_TYPES:
                    if t in type_stats:
                        ws.cell(row, 1, f'  {t}')
                        ws.cell(row, 2, f'{type_stats[t]} 项')
                        row += 1
                row += 1

        # 写入建议项统计
        sugg = self.data.get('suggestions', [])
        if sugg:
            ws.cell(row, 1, '建议项清单统计').font = Font(bold=True, size=12, color='0070C0')
            row += 1

            # 总计
            ws.cell(row, 1, '总计')
            ws.cell(row, 2, f'{len(sugg)} 项')
            ws.cell(row, 2).font = Font(bold=True, size=14, color='0070C0')
            row += 1

            # 按类别统计
            benefit_count = sum(1 for item in sugg if item.get('类别') == '加分项')
            optimize_count = sum(1 for item in sugg if item.get('类别') == '优化项')
            avoid_count = sum(1 for item in sugg if item.get('类别') == '避坑项')

            ws.cell(row, 1, '  加分项')
            ws.cell(row, 2, f'{benefit_count} 项')
            ws.cell(row, 2).font = Font(color='0070C0')
            row += 1

            ws.cell(row, 1, '  优化项')
            ws.cell(row, 2, f'{optimize_count} 项')
            ws.cell(row, 2).font = Font(color='00B050')
            row += 1

            ws.cell(row, 1, '  避坑项')
            ws.cell(row, 2, f'{avoid_count} 项')
            ws.cell(row, 2).font = Font(color='FF6600)

    def _create_disqualification_sheet(self):
        """创建废标风险检查表 Sheet（按类型分区分组）"""
        ws = self.wb.create_sheet('废标风险检查表')

        # 表头样式
        header_font = Font(size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # 类型分组标题样式
        group_font = Font(size=12, bold=True, color='FFFFFF')
        group_fill = PatternFill(start_color='7030A0', end_color='7030A0', fill_type='solid')
        group_alignment = Alignment(horizontal='left', vertical='center')

        # 边框样式
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 风险等级颜色
        risk_colors = {
            '高': 'C00000',
            '中': 'FF9900',
            '低': '008000'
        }

        # 设置列宽
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 35
        ws.column_dimensions['G'].width = 30
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 12
        ws.column_dimensions['J'].width = 20

        # 表头
        headers = ['序号', '风险等级', '类型', '检查项', '原文页码', '原文关键信息', '逻辑解释', '后果', '参考页码', '额外说明']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        ws.row_dimensions[1].height = 40

        # 按类型分组写入数据
        current_type = None
        display_row = 2

        for idx, item in enumerate(self.data.get('disqualification', []), 1):
            item_type = item.get('类型', '其他')

            # 如果类型变化，插入分组标题
            if item_type != current_type:
                current_type = item_type
                # 插入分组标题行
                ws.cell(display_row, 1, f'{current_type}（{sum(1 for x in self.data["disqualification"] if x.get("类型") == current_type)}项）')
                ws.cell(display_row, 1).font = group_font
                ws.cell(display_row, 1).fill = group_fill
                ws.cell(display_row, 1).alignment = group_alignment
                ws.merge_cells(f'A{display_row}:J{display_row}')
                ws.row_dimensions[display_row].height = 25
                display_row += 1

            risk_level = item.get('风险等级', '中')

            # 写入数据
            ws.cell(display_row, 1, idx)
            ws.cell(display_row, 2, risk_level)
            ws.cell(display_row, 3, item_type)
            ws.cell(display_row, 4, item.get('检查项', ''))
            ws.cell(display_row, 5, item.get('原文页码', ''))
            ws.cell(display_row, 6, item.get('原文关键信息', ''))
            ws.cell(display_row, 7, item.get('逻辑解释', ''))
            ws.cell(display_row, 8, item.get('后果', ''))
            ws.cell(display_row, 9, item.get('参考页码', ''))
            ws.cell(display_row, 10, item.get('额外说明', ''))

            # 设置风险等级颜色
            risk_color = risk_colors.get(risk_level, '000000')
            ws.cell(display_row, 2).font = Font(color=risk_color, bold=True)

            # 设置边框和对齐
            for col in range(1, 11):
                cell = ws.cell(display_row, col)
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

            ws.row_dimensions[display_row].height = 60
            display_row += 1

        # 冻结首行
        ws.freeze_panes = 'A2'

    def _create_suggestions_sheet(self):
        """创建建议项清单 Sheet"""
        ws = self.wb.create_sheet('建议项清单')

        # 表头样式
        header_font = Font(size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='008000', end_color='008000', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # 边框样式
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 类别颜色
        category_colors = {
            '加分项': '0070C0',
            '优化项': '00B050',
            '避坑项': 'FF6600'
        }

        # 检测是否有传统格式的建议项
        has_legacy = any(item.get('legacy') for item in self.data.get('suggestions', []))

        if has_legacy:
            # 传统格式表头
            headers = ['序号', '检查项', '类型', '原文页码', '原文关键信息', '逻辑解释', '参考页码', '额外说明']
        else:
            # 新格式表头
            headers = ['序号', '类别', '建议项', '分类', '对应条款/页码', '具体做法', '预期效果', '参考说明']

        # 设置列宽
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 10 if has_legacy else 12
        ws.column_dimensions['C'].width = 35 if has_legacy else 30
        ws.column_dimensions['D'].width = 12 if has_legacy else 12
        ws.column_dimensions['E'].width = 12 if has_legacy else 15
        ws.column_dimensions['F'].width = 40 if has_legacy else 35
        ws.column_dimensions['G'].width = 35 if has_legacy else 20
        ws.column_dimensions['H'].width = 15 if has_legacy else 20

        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        ws.row_dimensions[1].height = 30

        # 写入数据
        for idx, item in enumerate(self.data.get('suggestions', []), 2):
            if has_legacy:
                # 传统格式
                ws.cell(idx, 1, idx - 1)
                ws.cell(idx, 2, item.get('检查项', ''))
                ws.cell(idx, 3, item.get('类型', ''))
                ws.cell(idx, 4, item.get('原文页码', ''))
                ws.cell(idx, 5, item.get('原文关键信息', ''))
                ws.cell(idx, 6, item.get('逻辑解释', ''))
                ws.cell(idx, 7, item.get('参考页码', ''))
                ws.cell(idx, 8, item.get('额外说明', ''))
            else:
                # 新格式
                category = item.get('类别', '其他')
                ws.cell(idx, 1, idx - 1)
                ws.cell(idx, 2, category)
                ws.cell(idx, 3, item.get('建议项', ''))
                ws.cell(idx, 4, item.get('分类', ''))
                ws.cell(idx, 5, item.get('对应条款', ''))
                ws.cell(idx, 6, item.get('具体做法', ''))
                ws.cell(idx, 7, item.get('预期效果', ''))

                # 添加参考说明（加分潜力/优先级/风险等级）
                reference = item.get('参考条款', '')
                if category == '加分项':
                    ref = item.get('加分潜力', '')
                    ws.cell(idx, 8, f"{ref} | {reference}" if ref else reference)
                elif category == '优化项':
                    ref = item.get('优先级', '')
                    ws.cell(idx, 8, f"{ref} | {reference}" if ref else reference)
                elif category == '避坑项':
                    ref = item.get('风险等级', '')
                    ws.cell(idx, 8, f"{ref} | {reference}" if ref else reference)
                else:
                    ws.cell(idx, 8, reference)

                # 设置类别颜色
                category_color = category_colors.get(category, '000000')
                ws.cell(idx, 2).font = Font(color=category_color, bold=True)

            # 设置边框和对齐
            for col in range(1, len(headers) + 1):
                cell = ws.cell(idx, col)
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

            ws.row_dimensions[idx].height = 50

        # 冻结首行
        ws.freeze_panes = 'A2'

    def _create_validation_sheet(self):
        """创建 S 级验证报告 Sheet"""
        ws = self.wb.create_sheet('S级验证报告')

        # 标题样式
        title_font = Font(size=12, bold=True)
        header_fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')

        row = 1
        ws.cell(row, 1, 'S级准确性验证报告').font = Font(size=14, bold=True, color='4472C4')
        row += 2

        # 写入验证数据
        if self.data.get('validation'):
            validation = self.data['validation']

            # 识别统计
            if validation.get('statistics'):
                ws.cell(row, 1, '识别统计').font = title_font
                ws.cell(row, 1).fill = header_fill
                row += 1
                for key, value in validation['statistics'].items():
                    ws.cell(row, 1, key)
                    ws.cell(row, 2, value)
                    row += 1
                row += 1

            # 5轮验证
            if validation.get('rounds'):
                ws.cell(row, 1, '5轮验证执行情况').font = title_font
                ws.cell(row, 1).fill = header_fill
                row += 1
                for round_info in validation['rounds']:
                    ws.cell(row, 1, round_info.get('name', ''))
                    ws.cell(row, 2, round_info.get('status', ''))
                    row += 1
                    if round_info.get('items'):
                        for item in round_info['items']:
                            ws.cell(row, 1, '')
                            ws.cell(row, 2, f"  • {item}")
                            row += 1
                    row += 1

            # 质量得分
            if validation.get('scores'):
                ws.cell(row, 1, 'S级质量得分').font = title_font
                ws.cell(row, 1).fill = header_fill
                row += 1
                ws.cell(row, 1, '维度')
                ws.cell(row, 2, '得分')
                ws.cell(row, 3, '满分')
                ws.cell(row, 4, '说明')
                row += 1
                for score in validation['scores']:
                    ws.cell(row, 1, score.get('dimension', ''))
                    ws.cell(row, 2, score.get('score', ''))
                    ws.cell(row, 3, score.get('max_score', ''))
                    ws.cell(row, 4, score.get('note', ''))
                    row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 30


def main():
    """主函数"""
    import sys

    if len(sys.argv) < 2:
        print("招标文件检查清单 Excel 生成器 v4.1")
        print("")
        print("用法:")
        print("  python generate_excel.py <输入文件> [输出文件]")
        print("")
        print("示例:")
        print("  python generate_excel.py data/output/检查清单.md")
        print("  python generate_excel.py data/output/检查清单.md output.xlsx")
        print("")
        print("Excel 结构（根据PDF官方要求）:")
        print("  Sheet 1: 统计摘要（详细总数和分类统计）")
        print("  Sheet 2: 废标风险检查表（按类型分区分组）")
        print("  Sheet 3: 建议项清单（规范性建议要求）")
        print("  Sheet 4: S级验证报告")
        return

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    # 如果没有指定输出文件，自动生成
    if not output_file:
        input_path = Path(input_file)
        output_file = str(input_path.parent / f"{input_path.stem}.xlsx")

    # 创建生成器
    gen = ExcelGenerator()

    # 解析输入文件
    if input_file.endswith('.json'):
        success = gen.parse_json_data(input_file)
    else:
        success = gen.parse_markdown_file(input_file)

    if not success:
        print("解析输入文件失败")
        return

    # 生成 Excel
    gen.generate_excel(output_file)


if __name__ == '__main__':
    main()
